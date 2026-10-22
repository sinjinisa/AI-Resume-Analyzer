
from pathlib import Path
from docx import Document

# Find the folder containing this Python script
project_folder = Path(__file__).resolve().parent

# Create the sample_resumes folder if needed
resume_folder = project_folder / "sample_resumes"
resume_folder.mkdir(parents=True, exist_ok=True)

# Define the resume file location
resume_path = resume_folder / "demo_resume_v2.docx"

# Create the Word document
doc = Document()

doc.add_heading("Sample Resume", 0)
doc.add_paragraph("Skills: Python, SQL, Machine Learning")
doc.add_paragraph("Education: B.Tech Computer Science")
doc.add_paragraph("Projects: Student Data Analysis")

# Save the document
doc.save(str(resume_path))

print("Resume created successfully!")
print("Location:", resume_path)
print("File exists:", resume_path.is_file())
