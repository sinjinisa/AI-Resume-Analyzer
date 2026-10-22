
from pathlib import Path
from docx import Document
from resume_parser import extract_resume_text

# STEP 1: Create the sample_resumes folder
folder = Path("sample_resumes")
folder.mkdir(parents=True, exist_ok=True)

# STEP 2: Create a fictional resume
doc = Document()

doc.add_heading("Sample Resume", 0)
doc.add_paragraph("Skills: Python, SQL, Machine Learning")
doc.add_paragraph("Education: B.Tech Computer Science")
doc.add_paragraph("Projects: Student Data Analysis")

# STEP 3: Save the Word document
file_path = folder / "demo_resume.docx"
doc.save(file_path)

# STEP 4: Extract the resume text
text = extract_resume_text(file_path)

# STEP 5: Display the extracted text
print("EXTRACTED RESUME TEXT:")
print(text)
